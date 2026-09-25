from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from evaluation.synthetic_tabular import ROWS

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def margins(cell, top=80, start=80, bottom=80, end=80):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

doc=Document(); sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11); sec.page_height=Inches(8.5); sec.top_margin=Inches(.45); sec.bottom_margin=Inches(.45); sec.left_margin=Inches(.4); sec.right_margin=Inches(.4)
styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(7.5)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('SYNTHETIC PII EVALUATION DATASET'); r.bold=True; r.font.size=Pt(16); r.font.name='Arial'
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('100 tabular records modeled on the contact-information tables in the supplied Red Herring Prospectus'); r.italic=True; r.font.size=Pt(9)

p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
r=p.add_run('Evaluation note: '); r.bold=True; r.font.size=Pt(8)
r=p.add_run('Some rows intentionally use difficult formatting (initial-only names, unlabeled dates, Pvt. Ltd. variants and alternate address forms) so the holdout measures realistic false negatives and boundary errors.'); r.font.size=Pt(8)

headers=['RECORD','COMPANY','REGISTERED / MAILING ADDRESS','CONTACT PERSON','E-MAIL AND TELEPHONE','DATE OF BIRTH','SSN','CREDIT CARD','IP ADDRESS']
t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.autofit=False
widths=[.55,1.45,2.35,1.05,1.75,1.0,.85,1.15,.8]
for j,h in enumerate(headers):
    c=t.rows[0].cells[j]; c.width=Inches(widths[j]); shade(c,'C00000'); margins(c,100,80,100,80); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=p.add_run(h); rr.bold=True; rr.font.color.rgb=RGBColor(255,255,255); rr.font.size=Pt(7)
# repeat header
trPr=t.rows[0]._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)
for rdata in ROWS:
    vals=[rdata['record_id'],rdata['company'],rdata['address'],rdata['contact'],f"{rdata['email']}\n{rdata['phone']}",rdata['dob'],rdata['ssn'],rdata['card'],rdata['ip']]
    row=t.add_row()
    trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)
    cells=row.cells
    for j,val in enumerate(vals):
        cells[j].width=Inches(widths[j]); margins(cells[j],65,65,65,65); cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cells[j].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j in (0,3,5,6,7,8) else WD_ALIGN_PARAGRAPH.LEFT
        rr=p.add_run(val); rr.font.size=Pt(6.2)

out='evaluation/Synthetic_PII_Test_Dataset_100_Records.docx'; doc.save(out); print(out)
