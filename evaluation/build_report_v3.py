import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)
def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.autofit=False
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; shade(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
        if widths: c.width=Inches(widths[j])
    set_repeat(t.rows[0])
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row):
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[j].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(str(v)); r.font.size=Pt(9)
            if widths: cells[j].width=Inches(widths[j])
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

with open('evaluation/synthetic_results.json') as f: m=json.load(f)
with open('output/redaction_run.json') as f: run=json.load(f)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5)
for s in ['Heading 1','Heading 2']:
    styles[s].font.name='Aptos'; styles[s].font.color.rgb=RGBColor(31,78,120)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('PII Redaction Tool - Evaluation Report'); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor(31,78,120)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Red Herring Prospectus | Assignment Submission'); r.italic=True; r.font.size=Pt(11)

doc.add_heading('1. Evaluation methodology', level=1)
doc.add_paragraph('Evaluation is reported in three layers: (1) operational statistics from the complete supplied prospectus, (2) a small controlled regression benchmark that verifies coverage of all required PII types, and (3) a separate 100-record synthetic tabular holdout used for realistic precision, recall and F1 measurement. This avoids treating raw detections in an unannotated document as ground-truth accuracy.')
doc.add_paragraph('For the labeled holdout, a true positive requires both the PII category and exact detected text to match the annotation. Extra detections are false positives and missed annotations are false negatives. Cell accuracy measures whether an individual table cell is classified exactly; strict record accuracy requires every evaluated field in a record to match its ground truth.')

doc.add_heading('2. Evaluation on the supplied Red Herring Prospectus', level=1)
doc.add_paragraph('The deterministic offline pipeline was executed on the complete supplied DOCX. These are measured detection/replacement counts from the assignment document, not inferred accuracy values.')
counts=run.get('counts',run)
rows=[]
labels=[('COMPANY','Company entities'),('PERSON','Person entities'),('ADDRESS','Addresses'),('PHONE','Phone numbers'),('EMAIL','Email addresses'),('SSN','SSNs'),('CREDIT_CARD','Credit-card numbers'),('DOB','Dates of birth'),('IP_ADDRESS','IPv4 addresses')]
for k,n in labels: rows.append([n, counts.get(k,0)])
add_table(doc,['PII type','Detected / replaced'],rows,[4.5,1.7])
doc.add_paragraph('Because the full 127-page source was not exhaustively hand-annotated, precision and recall are not claimed for this section. Exact metrics are instead calculated on the labeled holdout below.')

doc.add_heading('3. Controlled regression benchmark', level=1)
doc.add_paragraph('The included baseline fixture contains known positive examples for every required PII category plus negative Order/Ticket examples. It achieved 9 TP, 0 FP and 0 FN (100% precision and recall). This is retained only as a deterministic regression/coverage check, not as a real-world performance claim.')

doc.add_page_break()
doc.add_heading('4. Synthetic tabular holdout - 100 records', level=1)
doc.add_paragraph('A separate 100-record DOCX dataset was generated in a table layout modeled on the contact-information tables in the supplied prospectus. Each record contains company, registered/mailing address, contact person, email, telephone, DOB, SSN, credit card and IP address fields. Difficult variants are intentionally included: initial-only and single-token names, unlabeled dates, alternate company suffixes and alternate address formatting.')
summary=[
 ['Records',m['records']],['Evaluated cells',m['cells']],['True positives',m['tp']],['False positives',m['fp']],['False negatives',m['fn']],
 ['Precision',f"{m['precision']*100:.2f}%"],['Recall',f"{m['recall']*100:.2f}%"],['F1 score',f"{m['f1']*100:.2f}%"],['Cell accuracy',f"{m['cell_accuracy']*100:.2f}%"],['Strict record accuracy',f"{m['record_accuracy']*100:.2f}%"]]
add_table(doc,['Metric','Measured value'],summary,[4.5,1.7])

doc.add_heading('5. Per-category performance', level=1)
prows=[]
for label,v in m['per_category'].items():
    prows.append([label,v['tp'],v['fp'],v['fn'],f"{v['precision']*100:.1f}%",f"{v['recall']*100:.1f}%",f"{v['f1']*100:.1f}%"])
add_table(doc,['PII type','TP','FP','FN','Precision','Recall','F1'],prows,[1.35,.55,.55,.55,1.0,1.0,1.0])

doc.add_heading('6. Interpretation and observed failure modes', level=1)
for txt in [
 'Structured patterns (email, phone, SSN, credit card and IPv4) were detected consistently in this holdout.',
 'Person recall decreases for initials and single-token names because the conservative fallback expects conventional multi-token names.',
 'DOB exact-match performance is lower because unlabeled dates are intentionally ambiguous and some valid contextual dates expose a boundary issue in the current regex.',
 'Company matching is sensitive to punctuation and suffix variants such as “Pvt. Ltd.”, which creates exact-match boundary errors.',
 'These failures are intentionally retained rather than tuning the benchmark to 100%; they provide concrete extension points for a production version.'
]: doc.add_paragraph(txt, style='List Bullet')

doc.add_heading('7. Reproducibility', level=1)
doc.add_paragraph('Synthetic dataset: evaluation/Synthetic_PII_Test_Dataset_100_Records.docx')
doc.add_paragraph('Synthetic evaluator: PYTHONPATH=. python evaluation/evaluate_synthetic.py')
doc.add_paragraph('Machine-readable results: evaluation/synthetic_results.json')
doc.add_paragraph('Prospectus run: python -m src.main input.docx output/Red_Herring_Prospectus_Redacted.docx --no-spacy --report output/redaction_run.json')

doc.save('Evaluation_Report.docx')
print('Evaluation_Report.docx')
